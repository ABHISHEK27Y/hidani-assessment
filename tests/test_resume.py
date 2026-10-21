import io

import fitz
import pytest
from docx import Document

from app.resume import extract_profile, read_resume
from app.skills import extract_skills


def test_work_dates_and_projects_do_not_count_education_as_experience():
    profile = extract_profile("""Sample Candidate
candidate@example.test
+91 9876543210
EDUCATION
Example Institute 2023 - Present
B.Tech Computer Science
WORK EXPERIENCE
Software Engineering Intern - Example Lab June 2026 - August 2026
Built a Python API and React interface.
PROJECTS
CloudCanvas - Visual infrastructure March 2026 - August 2026
Built an editor.
FraudWatch - Detection platform Jan 2026 - April 2026
Built a classifier.
ColdChain - IoT dashboard March 2025 - Sept 2025
Built a dashboard.
TECHNICAL SKILLS
Python TypeScript JavaScript C++ React.js Next.js Express.js TailwindCSS RAG Terraform
""")
    assert profile["designation"] == "Software Engineering Intern"
    assert profile["experience_months"] == 3
    assert profile["years_experience"] == .25
    assert len(profile["project_titles"]) == 3
    assert {"C++", "React", "Express.js", "Tailwind", "RAG", "Terraform"} <= set(profile["skills"])
    assert "ColdChain" in " ".join(profile["projects"])


def test_overlapping_work_dates_are_not_double_counted():
    profile = extract_profile("Sample Person\nWORK EXPERIENCE\nSoftware Engineer Jan 2025 - June 2025\nSoftware Developer March 2025 - August 2025")
    assert profile["experience_months"] == 8


def test_symbol_skills_and_word_boundaries():
    skills = extract_skills("C++ C# React.js PostgreSQL TailwindCSS; not a Spring season or Java island")
    assert {"C++", "C#", "React", "PostgreSQL", "Tailwind"} <= set(skills)
    assert "Spring" not in skills
    assert "Java" not in extract_skills("JavaScript")


def test_docx_table_order_and_text():
    document = Document()
    document.add_paragraph("Sample Candidate\ncandidate@example.test")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "TECHNICAL SKILLS\nPython React TypeScript"
    document.add_paragraph("WORK EXPERIENCE\nSoftware Engineer Jan 2025 - June 2025\nBuilt an accessible data dashboard and a backend API.")
    stream = io.BytesIO(); document.save(stream)
    text = read_resume("cv.docx", stream.getvalue())
    assert text.index("TECHNICAL SKILLS") < text.index("WORK EXPERIENCE")
    assert extract_profile(text)["experience_months"] == 6


@pytest.mark.parametrize("filename,data", [("cv.pdf", b"garbage"), ("cv.docx", b"garbage"), ("cv.exe", b"bad"), ("cv.pdf", b"")])
def test_reject_unreadable_files(filename, data):
    with pytest.raises(ValueError):
        read_resume(filename, data)


def test_reject_scan_large_and_password_pdf():
    document = fitz.open(); document.new_page()
    with pytest.raises(ValueError, match="readable text"):
        read_resume("scan.pdf", document.tobytes())
    locked = document.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="test-password")
    with pytest.raises(ValueError, match="password"):
        read_resume("locked.pdf", locked)
    with pytest.raises(ValueError, match="8 MB"):
        read_resume("large.pdf", b"x" * (8 * 1024 * 1024 + 1))
    document.close()
